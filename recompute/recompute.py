#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""recompute.py — 只用 VFB 上的資料，重算 Nern et al. 2025 的細胞型普查，並逐列跟論文比對。

論文：Nern et al., 2025, Nature 641(8065): 1225-1237
      Connectome-driven neural inventory of a complete visual system
      doi:10.1038/s41586-025-08746-0   CC-BY
      正文：We classified around 53,000 visual system neurons into 732 cell types

跑法：
    pip install openpyxl
    python3 recompute.py                # 約 40 秒，產出 out/ 底下四個檔案
    python3 make_figure.py              # 需要 matplotlib，畫逐型的散點圖

產出：
    out/target.json      論文那一側的靶（從補充表算出來）
    out/vfb.json         VFB 那一側（從知識庫取回）
    out/compare.json     逐列比對的結果
    out/fetched_at.txt   抓取時間與每一個關鍵數字

這支程式不需要 API key，也不需要登入。
"""
from __future__ import annotations

import collections
import html
import io
import json
import re
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path

# ══════════════════════════════════════════════════════════════════
# 自己決定的參數：全部集中在這裡。改動這一段就會改動結論。
# ══════════════════════════════════════════════════════════════════
PARAMS = {
    # 用 VFB 的哪一套資料。換一套，下面每個數字都會變。
    "vfb_dataset": "Nern2024",

    # 拿什麼當比對的鍵。兩個選擇實測命中率差很多：
    #   "instance" → 用神經元自己的標籤，取括號前那一段（例如 Dm20_R）
    #   "class"    → 用它所屬類別的標籤，取最後一個詞（例如 Tm3）
    # 預設用 instance；理由與兩者的命中率見 out/compare.json 的 key_choice。
    "match_key": "instance",

    # 「什麼算一個細胞型」：只算沒有下位類別的類別（末端類別）。
    # 設成 False 會把泛稱（例如「膽鹼性神經元」）也算成一個型，
    # 而那會讓型底下的顆數加總超過神經元總數。
    "leaf_classes_only": True,
}

KB = "https://pdb.virtualflybrain.org/db/neo4j/tx/commit"      # 唯讀，不需認證
SUPP = ("https://static-content.springer.com/esm/"
        "art%3A10.1038%2Fs41586-025-08746-0/MediaObjects/"
        "41586_2025_8746_MOESM4_ESM.zip")
SUPP_TABLE = "Sup_Table_1_Cell-types_and_counts_final.xlsx"
UA = "bsc-vfb-recompute/1.0"
OUT = Path(__file__).resolve().parent / "out"


# ══════════════════════════════════════════════════════════════════
def cypher(statement: str) -> list:
    """對 VFB 的知識庫下一句唯讀查詢。

    為什麼用知識庫而不用 REST：REST 的 run_query 回的是**列數**不是個數
    （同一顆神經元可以有好幾張影像），而且翻頁到某個位移就會停、不報錯。
    實測同一個資料集，REST 報 60,002 列並在 50,000 列截斷，知識庫算出 53,402 顆。
    """
    body = json.dumps({"statements": [{"statement": statement}]}).encode()
    req = urllib.request.Request(KB, data=body, method="POST",
                                 headers={"Content-Type": "application/json",
                                          "Accept": "application/json;charset=UTF-8",
                                          "User-Agent": UA})
    with urllib.request.urlopen(req, timeout=300) as r:
        d = json.loads(r.read().decode("utf-8", errors="replace"))
    if d.get("errors"):
        raise SystemExit(f"知識庫回錯誤：{d['errors']}")
    return [row["row"] for row in d["results"][0]["data"]]


def paper_target() -> dict:
    """論文那一側：下載補充表，算出要比對的靶。

    **下載網址有坑**：PMC 那條路回 200 但只有 1.8 KB，是空殼。
    要從出版社的靜態內容站抓。抓到的東西太小就直接停，不要往下做。
    """
    import openpyxl
    req = urllib.request.Request(SUPP, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=300) as r:
        blob = r.read()
    if len(blob) < 100_000:
        raise SystemExit(f"補充表只抓到 {len(blob)} 位元組，八成是空殼")
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        wb = openpyxl.load_workbook(io.BytesIO(z.open(SUPP_TABLE).read()), data_only=True)
    ws = wb[wb.sheetnames[0]]

    def n(x):
        try:
            return int(str(x).replace(",", "").strip())
        except Exception:
            return 0

    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]
    # bodyId 那一欄只給「圖裡用的那一顆」的編號，不是整列的每一顆；
    # 但拿來確認「這一列的神經元在不在 VFB 裡」已經夠——見 resolve_by_body_id()。
    by_instance = {r[1]: {"cell_type": r[0], "cells": n(r[2]), "group": r[3],
                          "body": n(r[4])} for r in rows}

    # 列數比型數多，因為**一列是一個（型, 側別）實例**，不是一個型。
    # 論文正文報的是型；表列的是列。不分清楚，778 與 732 會看起來像矛盾。
    per_type = collections.defaultdict(list)
    for inst, v in by_instance.items():
        per_type[v["cell_type"]].append(inst)
    def side(x):
        return x[-2:] if x[-2:] in ("_L", "_R") else "其他"
    buckets = collections.Counter()
    groups = collections.defaultdict(collections.Counter)
    for ct, insts in per_type.items():
        s = tuple(sorted({side(i) for i in insts}))
        key = {("_L", "_R"): "both", ("_R",): "right_only", ("_L",): "left_only"}.get(s, str(s))
        buckets[key] += 1
        groups[key][by_instance[insts[0]]["group"]] += 1
    # 每一種型各貢獻幾列。**這一段是為了把「型」與「列」兜起來**——
    # 兩組數字並排放著而不寫出換算，讀的人只會看到 681 跟 635 打架。
    contributed = {}
    for key, r, l in (("right_only", 1, 0), ("both", 1, 1), ("left_only", 0, 1)):
        k = buckets[key]                     # 不要用 n：上面那個 n() 是轉數字的函式
        contributed[key] = {"_R": k * r, "_L": k * l, "rows": k * (r + l)}
    contributed["total"] = {
        k: sum(v[k] for v in contributed.values()) for k in ("_R", "_L", "rows")}

    sides = {"rows_per_type": dict(sorted(collections.Counter(
                 len(v) for v in per_type.values()).items())),
             "types_by_side": dict(buckets),
             "rows_contributed_by_side": contributed,
             "groups_by_side": {k: dict(v.most_common()) for k, v in groups.items()},
             "note": "778 列 − 732 型 = 46，剛好就是佔兩列的型數；那 46 個全部是 _L／_R 成對。"}


    return {"rows": len(rows),
            "distinct_cell_types": len({r[0] for r in rows}),
            "total_cells": sum(n(r[2]) for r in rows),
            "instance_side": dict(collections.Counter(side(i) for i in by_instance)),
            "sides": sides,
            "by_instance": by_instance}


def vfb_side() -> dict:
    """VFB 那一側：取回資料集底下每一顆神經元，以及類別的上下位關係。"""
    ds = PARAMS["vfb_dataset"]
    neurons = cypher(
        "MATCH (n:Individual)-[:has_source]->(:DataSet {short_form:'%s'}) "
        "RETURN n.short_form, n.label" % ds)

    # 每一顆屬於哪些類別
    pairs = cypher(
        "MATCH (n:Individual)-[:has_source]->(:DataSet {short_form:'%s'}) "
        "MATCH (n)-[:INSTANCEOF]->(c:Class) "
        "RETURN n.short_form, c.short_form, c.label" % ds)
    classes = {c: lab for _, c, lab in pairs}

    # 哪些類別是「泛稱」（底下還有別的類別，而那個別的類別也在這一組裡）
    lit = "[" + ",".join(f"'{c}'" for c in classes) + "]"
    # **只看這一組之內**：某個類別在整個 VFB 本體論裡還有下位，但那些下位
    # 在這套資料裡沒有神經元，它在這裡仍算末端。這是我們的選擇，不是資料說的。
    generic_rows = cypher(
        "MATCH (a:Class)<-[:SUBCLASSOF*1..8]-(b:Class) "
        "WHERE a.short_form IN %s AND b.short_form IN %s "
        "RETURN a.short_form, a.label, count(DISTINCT b) "
        "ORDER BY count(DISTINCT b) DESC" % (lit, lit))
    generic = {r[0] for r in generic_rows}
    leaf = {c: lab for c, lab in classes.items() if c not in generic}

    # 泛稱底下各自直接掛了幾顆——拿來說明「為什麼加總會超過總數」
    direct = dict(collections.Counter(
        c for _, c, _ in pairs if c in generic))
    generic_list = [{"short_form": sf, "label": lab, "subclasses_here": k,
                     "neurons_directly_here": direct.get(sf, 0)}
                    for sf, lab, k in generic_rows]

    # bodyId → 標籤。VFB 的標籤長這樣：Dm15_R (JRC_OpticLobe:65558)
    body = re.compile(r"\(JRC_OpticLobe:(\d+)\)")
    by_body = {}
    for _, lab in neurons:
        m = body.search(lab or "")
        if m:
            by_body[int(m.group(1))] = (lab or "").split(" (")[0]

    # 比對的鍵
    pre = re.compile(r"^(.*?)\s*\(")
    key_of_instance = collections.Counter()
    for _, lab in neurons:
        m = pre.match(lab or "")
        key_of_instance[m.group(1) if m else (lab or "")] += 1
    by_class = collections.Counter()
    for nid, c, lab in pairs:
        if PARAMS["leaf_classes_only"] and c in generic:
            continue
        by_class[(lab or "").split()[-1] if lab else ""] += 1

    return {"dataset": ds,
            "neurons": len(neurons),
            "by_body_id": by_body,
            "classes_all": len(classes),
            "classes_generic": len(generic),
            "classes_leaf": len(leaf),
            "generic_classes": generic_list,
            # 821 個末端類別的完整標籤。存下來，頁面上引到的任何一個型名
            # 都對得回這裡（稽核只認 out/ 裡出現過的字串）。
            "leaf_classes": {c: lab for c, lab in sorted(leaf.items())},
            "instanceof_pairs": len(pairs),
            "keys_instance": dict(key_of_instance),
            "keys_class": dict(by_class)}


def resolve_by_body_id(target: dict, vfb: dict, only_paper: list,
                       only_vfb: list) -> dict:
    """只在論文那一側的列，改用 bodyId 再查一次。

    **名字對不上不代表東西不在。**兩邊給同一顆重建的神經元取不同的型名，
    用名字比就永遠對不上；而 bodyId 是同一份重建的編號，兩邊共用。
    論文補充表有一欄 bodyId，VFB 的標籤裡也帶著它（`Dm15_R (JRC_OpticLobe:65558)`），
    所以這一步可以把「改名」跟「真的沒收進來」分開。

    回傳兩桶：
      renamed    bodyId 在 VFB 找得到，但 VFB 給它的名字不一樣
      absent     bodyId 在 VFB 的 Nern2024 裡完全不存在
    """
    paper, by_body = target["by_instance"], vfb["by_body_id"]
    got = vfb["keys_instance"]
    renamed, absent = {}, {}
    for inst in only_paper:
        b = paper[inst].get("body")
        lab = by_body.get(b) if b else None
        if lab:
            renamed[inst] = {"body": b, "vfb_name": lab,
                             "paper_cells": paper[inst]["cells"],
                             "vfb_cells": got.get(lab)}
        else:
            absent[inst] = {"body": b, "paper_cells": paper[inst]["cells"]}

    # 改名之後，VFB 那一側的名字就不該再算成「只在 VFB」
    merged = collections.Counter(v["vfb_name"] for v in renamed.values())

    # ── 把全部 778 列都對一次，可以直接驗論文〈Versions of the dataset〉那一段 ──
    # VFB 收的是 optic-lobe:v1.0.1，論文補充表是 v1.1。論文說 v1.1 相對 v1.0.1
    # 「5 個型被拆成 11 個新型」——這裡用編號獨立還原一次，看數字對不對得上。
    all_map, absent_all = {}, []
    for inst, d in paper.items():
        lab = by_body.get(d.get("body"))
        (all_map.__setitem__(inst, lab) if lab else absent_all.append(inst))
    back = collections.defaultdict(list)
    for inst, lab in all_map.items():
        back[lab].append(inst)
    splits = {k: sorted(v) for k, v in back.items() if len(v) > 1}
    version = {
        "vfb_release": "optic-lobe:v1.0.1（VFB 自己的資料集描述就這樣寫）",
        "paper_table_release": "optic-lobe:v1.1（論文正文與資料可用性都指這一版）",
        "paper_says": "5 types were split into a total of 11 new types "
                      "and 2 types were merged into a single type",
        "splits_recovered": splits,
        "n_v101_types_split": len(splits),
        "n_v11_types_produced": sum(len(v) for v in splits.values()),
        "rows_not_in_v101": sorted(absent_all),
        "groups_of_rows_not_in_v101": dict(collections.Counter(
            paper[i]["group"] for i in absent_all)),
        "caveat": "補充表的 bodyId 欄只給「圖裡用的那一顆」，所以這是型層級的對照，"
                  "不是逐顆對照。論文說的『2 個型併成 1 個』用這個方法驗不出來——"
                  "被併掉的那一個在 v1.1 沒有自己的列，沒有代表編號可查。",
    }
    return {"renamed": renamed, "absent": absent,
            "n_renamed": len(renamed), "n_absent": len(absent),
            "cells_absent": sum(v["paper_cells"] for v in absent.values()),
            # 論文拆成兩型、VFB 併成一型的，會在這裡露出來（同一個 vfb_name 出現兩次）
            "vfb_names_taking_more_than_one_paper_row":
                {k: v for k, v in merged.items() if v > 1},
            # 這些 VFB 名字其實就是論文那幾列，不該再算成「只在 VFB」
            "version_gap": version,
            "vfb_names_no_longer_only_in_vfb":
                sorted({v["vfb_name"] for v in renamed.values()} & set(only_vfb)),
            "only_in_vfb_after": len(set(only_vfb)
                - {v["vfb_name"] for v in renamed.values()}),
            # 分類：一對一改名／側別標到另一邊／論文拆而 VFB 併成一個
            "kinds": {
                "renamed_one_to_one": sorted(
                    k for k, v in renamed.items()
                    if merged[v["vfb_name"]] == 1
                    and v["paper_cells"] == v["vfb_cells"]),
                "side_label_differs": sorted(
                    k for k, v in renamed.items()
                    if merged[v["vfb_name"]] == 1
                    and v["paper_cells"] != v["vfb_cells"]),
                "paper_split_vfb_merged": {
                    name: sorted(k for k, v in renamed.items()
                                 if v["vfb_name"] == name)
                    for name, c in merged.items() if c > 1}}}


def class_tree_examples() -> dict:
    """PART 8 那張樹狀圖用到的實例——說明「末端類別」與「泛稱」的差別。

    **這一支存在的理由是一個被查出來是錯的直覺。**
    直覺：把泛稱也當成型，細胞會被數兩次（一次算在 Dm3a，一次算在 Dm3）。
    實際查：**Dm3 與 Dm3a／Dm3b 沒有任何一顆神經元同時掛在兩邊**（重疊 0）。
    真正造成重複的是**橫切的屬性**——4,063 顆神經元既有自己的細胞型，
    又掛在「膽鹼性神經元」底下。兩種泛稱要分開講。
    """
    ds = PARAMS["vfb_dataset"]

    def kids(sf):
        return [r[1] for r in cypher(
            "MATCH (:Class {short_form:'%s'})<-[:SUBCLASSOF]-(b:Class) "
            "RETURN b.short_form, b.label" % sf)]

    def cells(sf):
        return cypher(
            "MATCH (n:Individual)-[:has_source]->(:DataSet {short_form:'%s'}) "
            "MATCH (n)-[:INSTANCEOF]->(:Class {short_form:'%s'}) "
            "RETURN count(n)" % (ds, sf))[0][0]

    DM3, DM3A, DM3B, DM3C = ("FBbt_00003770", "FBbt_00052303",
                             "FBbt_00052304", "FBbt_00053403")
    DM15, LC10, CM11 = "FBbt_00111276", "FBbt_00100482", "FBbt_20007247"
    NT = ["FBbt_00058205", "FBbt_00058207", "FBbt_00058208"]   # 三種傳導物質泛稱

    # PART 8 那張樹上的每一個節點，逐個套一次規則。
    # **重點**：844 個不是本體論的所有節點，是「至少有一顆神經元被直接標成它」
    # 的類別。樹上有兩個節點一顆都沒有，所以根本不參與計數。
    in_set = {c for c, in [(r[0],) for r in cypher(
        "MATCH (n:Individual)-[:has_source]->(:DataSet {short_form:'%s'}) "
        "MATCH (n)-[:INSTANCEOF]->(c:Class) RETURN DISTINCT c.short_form" % ds)]}
    tree_nodes = []
    for label, sf in (("adult neuron", "FBbt_00047095"),
                      ("adult optic lobe intrinsic neuron", "FBbt_00007577"),
                      ("distal medullary amacrine neuron", "FBbt_00003767"),
                      ("Dm15", DM15), ("Dm3", DM3),
                      ("Dm3a", DM3A), ("Dm3b", DM3B), ("Dm3c", DM3C)):
        here = cells(sf)
        # **兩個數字都要**：掛在樹上的節點旁邊只寫一個數字，讀者一定讀成子樹總數
        #   （Dm3 401 顆底下站著 Dm3a 601 顆，看起來就是矛盾）。
        # cells_directly_here 是「只被標到這麼細」的那些，cells_in_subtree 是
        # 「這一叢總共幾顆」——DISTINCT 是必要的，一顆神經元可以掛好幾個類別。
        subtree = cypher(
            "MATCH (n:Individual)-[:has_source]->(:DataSet {short_form:'%s'}) "
            "MATCH (n)-[:INSTANCEOF]->(:Class)-[:SUBCLASSOF*0..8]->"
            "(:Class {short_form:'%s'}) "
            "RETURN count(DISTINCT n)" % (ds, sf))[0][0]
        below = [r[0] for r in cypher(
            "MATCH (:Class {short_form:'%s'})<-[:SUBCLASSOF*1..8]-(b:Class) "
            "RETURN DISTINCT b.short_form" % sf)]
        below_here = [b for b in below if b in in_set]
        tree_nodes.append({
            "label": label, "short_form": sf,
            "cells_directly_here": here,
            "cells_in_subtree": subtree,
            "in_the_844": sf in in_set,
            "subclasses_within_the_844": len(below_here),
            "verdict": ("泛稱" if (sf in in_set and below_here)
                        else "末端類別" if sf in in_set
                        else "不在這 844 個裡（沒有神經元直接標成它）")})

    return {
        "what_the_844_are": "有至少一顆神經元被直接標成它的類別；"
                            "不是本體論裡的所有節點。",
        "tree_nodes": tree_nodes,
        "Dm3": kids(DM3), "LC10": kids(LC10), "Cm11": kids(CM11),
        "Dm15_has_no_subclass": not kids(DM15),
        "Dm15_parents": [r[1] for r in cypher(
            "MATCH (:Class {short_form:'%s'})-[:SUBCLASSOF]->(b:Class) "
            "RETURN b.short_form, b.label" % DM15)],
        "cells_in_this_dataset": {"Dm3": cells(DM3), "Dm3a": cells(DM3A),
                                  "Dm3b": cells(DM3B), "Dm3c": cells(DM3C),
                                  "Dm15": cells(DM15)},
        "overlap_Dm3_with_Dm3ab": cypher(
            "MATCH (n:Individual)-[:has_source]->(:DataSet {short_form:'%s'}) "
            "MATCH (n)-[:INSTANCEOF]->(:Class {short_form:'%s'}) "
            "MATCH (n)-[:INSTANCEOF]->(c:Class) WHERE c.short_form IN ['%s','%s'] "
            "RETURN count(DISTINCT n)" % (ds, DM3, DM3A, DM3B))[0][0],
        "neurons_with_both_a_leaf_type_and_a_transmitter_generic": {
            r[0]: r[1] for r in cypher(
                "MATCH (n:Individual)-[:has_source]->(:DataSet {short_form:'%s'}) "
                "MATCH (n)-[:INSTANCEOF]->(g:Class), (n)-[:INSTANCEOF]->(l:Class) "
                "WHERE g.short_form IN ['%s'] AND g<>l "
                "RETURN g.label, count(DISTINCT n) "
                "ORDER BY count(DISTINCT n) DESC" % (ds, "','".join(NT)))},
        "example_neuron_with_two_classes": cypher(
            "MATCH (n:Individual)-[:has_source]->(:DataSet {short_form:'%s'}) "
            "WHERE n.label STARTS WITH 'Mi10_R' WITH n LIMIT 1 "
            "MATCH (n)-[:INSTANCEOF]->(c:Class) "
            "RETURN n.label, collect(c.label)" % ds)[0],
        "note": "分類學上位（Dm3 之於 Dm3a／Dm3b）**不會**讓細胞被數兩次——重疊是 0；"
                "會重複的是橫切的屬性（神經傳導物質）。兩種泛稱要分開講。",
    }


def compare(target: dict, vfb: dict) -> dict:
    """逐列比對。**兩個鍵都算一次**，因為命中率的差距本身是結果的一部分。"""
    paper = target["by_instance"]
    out = {"key_choice": {}}
    for key in ("instance", "class"):
        got = vfb["keys_instance"] if key == "instance" else vfb["keys_class"]
        if key == "instance":
            hit = set(paper) & set(got)
            denom = len(paper)
        else:
            # 類別那一組沒有左右，所以拿論文的 cell type 欄比
            ptypes = {v["cell_type"] for v in paper.values()}
            hit = ptypes & set(got)
            denom = len(ptypes)
        out["key_choice"][key] = {"matched": len(hit), "of": denom,
                                  "rate_pct": round(len(hit) / denom * 100, 1),
                                  "denominator_is": ("論文表的 instance 列（帶側別）"
                                                     if key == "instance"
                                                     else "論文表的 cell type（不帶側別）")}

    # **上面兩個比率的分母不一樣**，因為兩種鍵的粒度不同：
    # VFB 的個體標籤帶側別（Dm20_R），類別標籤不帶（Dm15）。
    # 各自只能跟論文表對應的那一欄比。
    # 並排放兩個分母不同的百分比會誤導，所以再算一次同分母的版本：
    # 兩邊都降到「型」這一層，分母統一用論文的 732 個 cell type。
    ptypes = {v["cell_type"] for v in target["by_instance"].values()}
    side = re.compile(r"_[LR]$")
    inst_as_types = {side.sub("", k) for k in vfb["keys_instance"]}
    out["key_choice_same_denominator"] = {
        "denominator": len(ptypes),
        "denominator_is": "論文的 cell type 數，兩種鍵都降到型這一層",
        "instance": len(ptypes & inst_as_types),
        "class": len(ptypes & set(vfb["keys_class"])),
        "instance_pct": round(len(ptypes & inst_as_types) / len(ptypes) * 100, 1),
        "class_pct": round(len(ptypes & set(vfb["keys_class"])) / len(ptypes) * 100, 1),
        "gap_pct": round((len(ptypes & inst_as_types)
                          - len(ptypes & set(vfb["keys_class"]))) / len(ptypes) * 100, 1),
        "note": "降到型這一層會讓 instance 那一側少掉左右的區分，"
                "所以它的分子跟上面那個 753 不一樣——那 753 數的是列。"}

    # 這裡才是 PARAMS["match_key"] 真正生效的地方。
    # 選 class 的時候，論文那一側也要換成 cell type 欄，否則比的是兩種東西。
    if PARAMS["match_key"] == "instance":
        got = vfb["keys_instance"]
    else:
        got = vfb["keys_class"]
        paper = {}
        for v in target["by_instance"].values():
            paper.setdefault(v["cell_type"], {"cells": 0})["cells"] += v["cells"]
    both = sorted(set(paper) & set(got))
    diff = [{"instance": k, "paper": paper[k]["cells"], "vfb": got[k],
             "delta": got[k] - paper[k]["cells"]}
            for k in both if got[k] != paper[k]["cells"]]
    only_paper = sorted(set(paper) - set(got))
    only_vfb = sorted(set(got) - set(paper))

    # 論文把一個型拆成 a/b/c，而 VFB 還是合在一起的 —— 這一類可以自動認出來
    def base(x):
        m = re.match(r"^(.*?)([a-z])(_[LR])$", x)
        return (m.group(1) + m.group(3)) if m else None

    split = collections.defaultdict(list)
    for x in only_paper:
        b = base(x)
        if b and b in only_vfb:
            split[b].append(x)

    # 論文自己點名的那一群：R7／R8 光受器（正文圖說說它們被少算了）
    photo = {k: {"paper": paper[k]["cells"], "vfb": got.get(k)}
             for k in sorted(paper) if k.startswith(("R7", "R8"))}
    photo_delta = sum(v["vfb"] - v["paper"] for v in photo.values() if v["vfb"])

    prefix = collections.Counter()
    for k in only_vfb:
        m = re.match(r"^([A-Za-z]+)", k)
        if m:
            prefix[m.group(1)] += 1

    # bodyId 那一步是**逐列**的（論文補充表的 bodyId 欄一列一個代表編號），
    # 所以只有 match_key="instance" 時它才有對應的東西可查。
    # 換成 class 的時候 only_paper 裝的是型名，不是列名——不硬做，明講跳過。
    by_body_id = (resolve_by_body_id(target, vfb, only_paper, only_vfb)
                  if PARAMS["match_key"] == "instance" else
                  {"skipped": "只有 match_key='instance' 時才做：論文補充表的 "
                              "bodyId 是逐列（逐個型, 側別）給的，"
                              "換成 class 這一側就沒有可對的編號。"})

    d = [abs(x["delta"]) for x in diff]
    return {**out,
            "only_in_vfb_cells": sum(got[k] for k in only_vfb),
            "only_in_paper_explained_by_split": sum(len(v) for v in split.values()),
            "only_in_paper_unexplained": len(only_paper) - sum(len(v) for v in split.values()),
            "total_cells_delta": vfb["neurons"] - target["total_cells"],
            "photoreceptor_rows": photo,
            "photoreceptor_delta_total": photo_delta,
            "only_in_vfb_prefix_top": dict(prefix.most_common(8)),
            "matched": len(both),
            "same_count": len(both) - len(diff),
            "different_count": len(diff),
            "differences": sorted(diff, key=lambda x: -abs(x["delta"])),
            "delta_max": max(d) if d else 0,
            "delta_is_one": sum(1 for x in d if x == 1),
            "delta_over_ten": sum(1 for x in d if x > 10),
            "only_in_paper": only_paper,
            "only_in_vfb": only_vfb,
            "paper_split_vfb_did_not": {k: sorted(v) for k, v in split.items()},
            "by_body_id": by_body_id,
            "total_cells_paper": target["total_cells"],
            "total_cells_vfb": vfb["neurons"]}


# 頁面上引用的論文原句。**每一句都要逐字對回全文**——
# paper_quotes() 每次重跑都會重驗，對不上就停在那裡，不會悄悄產出一份假的。
PAPER_QUOTES = {
    "n_types": "We classified around 53,000 visual system neurons into 732 cell types",
    "versions": "Two releases of the optic lobe dataset are available: the initial "
                "release (optic-lobe:v1.0.1) and the current version of the dataset "
                "(optic-lobe:v1.1).",
    "split_merge": "We also revised some cell-type definitions (5 types were split "
                   "into a total of 11 new types and 2 types were merged into a "
                   "single type).",
    "completeness": "optic-lobe:v1.1 incorporates updates to the segmentation "
                    "(mainly some merges of small, previously unconnected parts of "
                    "reconstructed segments with these neurons); these changes "
                    "collectively result in a small but detectable increase of "
                    "reconstruction completeness",
    "retyping": "adjusted the typing of some individual cells (nearly all of these "
                "are R7 and R8 photoreceptors",
    "tbd": "In three cases, we applied a placeholder name, adding \u2018_TBD\u2019, "
           "because we anticipate that future central brain data will provide a "
           "more appropriate designation.",
}

EPMC = ("https://www.ebi.ac.uk/europepmc/webservices/rest/"
        "PMC12119369/fullTextXML")


def paper_quotes() -> dict:
    """把頁面上引的論文原句逐條對回全文。

    **為什麼走 Europe PMC**：PMC 的網頁版擋自動存取（回 reCAPTCHA 頁），
    而 fullTextXML 這條給的是全文。這篇是 CC-BY。

    **為什麼要有這一支**：這個檔案原本是臨時腳本產的，
    於是「逐字比對過」這句宣稱沒有任何程式在維持——
    乾淨重跑不會重驗，論文換版或引錯字也不會有人發現。現在對不上就直接停。
    """
    req = urllib.request.Request(EPMC, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=120) as r:
        xml = r.read().decode("utf-8", errors="replace")
    # 去標籤 → 還原實體 → 把換行壓成單一空白，才對得上跨行的句子
    text = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", xml)))

    missing = [k for k, v in PAPER_QUOTES.items() if v not in text]
    if missing:
        raise SystemExit(f"這幾句在全文裡逐字對不到，不要往下做：{missing}")
    return {
        "source": "Nern et al. 2025 全文，經 Europe PMC 的 fullTextXML 取得（PMC12119369）",
        "url": EPMC,
        "pmcid": "PMC12119369",
        "fetched": datetime.now(timezone.utc).date().isoformat(),
        "check": f"{len(PAPER_QUOTES)} 條引文在上面那份全文裡逐字對得到"
                 "（本檔每次重新產生時都重驗一次，對不上就中止）",
        "full_text_chars": len(text),
        **PAPER_QUOTES,
    }


def main() -> int:
    OUT.mkdir(exist_ok=True)
    print("① 取論文的靶…")
    target = paper_target()
    print(f"   {target['rows']} 列　{target['distinct_cell_types']} 個型　"
          f"{target['total_cells']:,} 顆")

    print("② 取 VFB 那一側…")
    vfb = vfb_side()
    print(f"   {vfb['neurons']:,} 顆　類別 {vfb['classes_all']}"
          f"（泛稱 {vfb['classes_generic']}、末端 {vfb['classes_leaf']}）")

    print("③ 逐列比對…")
    cmp = compare(target, vfb)
    for k, v in cmp["key_choice"].items():
        print(f"   鍵={k:9s} 對得上 {v['matched']}/{v['of']}　{v['rate_pct']}%"
              f"　（分母：{v['denominator_is']}）")
    d = cmp["key_choice_same_denominator"]
    print(f"   放到同一個分母（{d['denominator']} 個型）："
          f"instance {d['instance_pct']}% vs class {d['class_pct']}%"
          f"　差 {d['gap_pct']} 個百分點")
    print(f"   顆數相同 {cmp['same_count']}　不同 {cmp['different_count']}"
          f"（最大差 {cmp['delta_max']}）")
    print(f"   只在論文 {len(cmp['only_in_paper'])}　只在 VFB {len(cmp['only_in_vfb'])}")

    print("④ 取樹狀圖的實例…")
    tree = class_tree_examples()
    print(f"   Dm3 與 Dm3a／Dm3b 的重疊：{tree['overlap_Dm3_with_Dm3ab']} 顆"
          f"（分類學上位不會重複數）")
    _dm3 = next(n for n in tree["tree_nodes"] if n["label"] == "Dm3")
    print(f"   Dm3 直接 {_dm3['cells_directly_here']} 顆 ／ 這一叢共 "
          f"{_dm3['cells_in_subtree']} 顆（樹狀圖上兩個都要寫）")

    print("⑤ 對論文原句…")
    quotes = paper_quotes()
    print(f"   {len(PAPER_QUOTES)} 條全部逐字對得到")

    for name, obj in (("target", target), ("vfb", vfb), ("compare", cmp),
                      ("class_tree_examples", tree), ("paper_quotes", quotes)):
        (OUT / f"{name}.json").write_text(
            json.dumps(obj, ensure_ascii=False, indent=1) + "\n", "utf-8")
    (OUT / "fetched_at.txt").write_text(
        f"抓取時間 {datetime.now(timezone.utc).isoformat(timespec='seconds')}（UTC）\n"
        f"參數 {json.dumps(PARAMS, ensure_ascii=False)}\n"
        f"論文 {target['distinct_cell_types']} 型 / {target['total_cells']} 顆\n"
        f"VFB  {vfb['neurons']} 顆 / 末端類別 {vfb['classes_leaf']}\n"
        f"對得上 {cmp['matched']}　顆數相同 {cmp['same_count']}\n", "utf-8")
    print(f"→ {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
